import os
import shutil
import time
import uuid
from html import escape
from typing import Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, HTMLResponse
from sqlalchemy.orm import Session

from app.database import get_db
from app.models.report import Report
from app.services.report_service import ReportService

router = APIRouter(prefix="/api/reports", tags=["reports"])


def _delete_file_with_retry(file_path: str, retries: int = 5, delay_seconds: float = 0.25) -> None:
    """Retry file deletion to handle brief Windows Excel lock periods."""
    normalized_path = os.path.normpath(file_path)
    if not os.path.isabs(normalized_path):
        normalized_path = os.path.abspath(normalized_path)

    last_error: Optional[Exception] = None
    for attempt in range(retries):
        try:
            if os.path.exists(normalized_path):
                os.remove(normalized_path)
            return
        except FileNotFoundError:
            return
        except PermissionError as exc:
            last_error = exc
            if attempt < retries - 1:
                time.sleep(delay_seconds * (attempt + 1))
                continue
            raise

    if last_error is not None:
        raise last_error


@router.post("/upload")
async def upload_and_generate(
    eod_file: Optional[UploadFile] = File(None),
    morning_file: Optional[UploadFile] = File(None),
    snapshot_type: str = Form("EOD"),          # "MORNING" | "EOD"
    report_date: Optional[str] = Form(None),
    day1_fixed_scope: Optional[float] = Form(None),
    sprint_start: Optional[str] = Form(None),
    sprint_end: Optional[str] = Form(None),
    db: Session = Depends(get_db),
):
    """
    Upload CSVs and generate the sprint report.

    Scenario 1 — MORNING only : snapshot_type=MORNING, morning_file=<file>
    Scenario 2 — EOD only     : snapshot_type=EOD,     eod_file=<file>
    Scenario 3 — Both         : snapshot_type=EOD,     eod_file=<file>, morning_file=<file>
    """
    os.makedirs("./uploads", exist_ok=True)

    # Validate: at least one file must be supplied
    morning_provided = morning_file and morning_file.filename
    eod_provided = eod_file and eod_file.filename

    if not morning_provided and not eod_provided:
        raise HTTPException(status_code=422, detail="At least one CSV file (Morning or EOD) is required.")

    # Scenario 1: morning-only — swap so we process morning as the primary file
    if snapshot_type == "MORNING":
        if not morning_provided:
            raise HTTPException(status_code=422, detail="Morning CSV is required for MORNING snapshot type.")
        primary_file = morning_file
        secondary_file = None
        original_filename = morning_file.filename
    else:
        if not eod_provided:
            raise HTTPException(status_code=422, detail="EOD CSV is required for EOD snapshot type.")
        primary_file = eod_file
        secondary_file = morning_file if morning_provided else None
        original_filename = eod_file.filename

    # Save primary file
    primary_ext = os.path.splitext(primary_file.filename)[1] or ".csv"
    # Support .csv, .tsv, .txt
    if primary_ext.lower() not in (".csv", ".tsv", ".txt"):
        primary_ext = ".csv"
    primary_path = f"./uploads/{uuid.uuid4().hex}{primary_ext}"
    with open(primary_path, "wb") as f:
        shutil.copyfileobj(primary_file.file, f)

    # Save secondary (morning) file if provided
    secondary_path = None
    if secondary_file:
        sec_ext = os.path.splitext(secondary_file.filename)[1] or ".csv"
        if sec_ext.lower() not in (".csv", ".tsv", ".txt"):
            sec_ext = ".csv"
        secondary_path = f"./uploads/{uuid.uuid4().hex}{sec_ext}"
        with open(secondary_path, "wb") as f:
            shutil.copyfileobj(secondary_file.file, f)

    try:
        service = ReportService(db)
        result = service.process_upload(
            file_path=primary_path,
            snapshot_type=snapshot_type,
            morning_file_path=secondary_path,   # only set for Scenario 3
            report_date=report_date,
            day1_fixed_scope=day1_fixed_scope,
            sprint_start=sprint_start,
            sprint_end=sprint_end,
            original_filename=original_filename,
        )
        return result
    except Exception as exc:
        raise HTTPException(status_code=422, detail=str(exc))
    finally:
        if os.path.exists(primary_path):
            os.remove(primary_path)
        if secondary_path and os.path.exists(secondary_path):
            os.remove(secondary_path)


@router.post("/validate")
async def validate_csv(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    """Validate a CSV without generating a report."""
    os.makedirs("./uploads", exist_ok=True)
    ext = os.path.splitext(file.filename)[1]
    tmp_path = f"./uploads/validate_{uuid.uuid4().hex}{ext}"
    with open(tmp_path, "wb") as f:
        shutil.copyfileobj(file.file, f)
    try:
        service = ReportService(db)
        return service.validate_csv(tmp_path)
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


@router.get("/morning-snapshots")
def list_morning_snapshots(db: Session = Depends(get_db)):
    """List all saved morning snapshots."""
    service = ReportService(db)
    return service.get_morning_snapshots()


@router.delete("/morning-snapshots/{snapshot_id}")
def delete_morning_snapshot(snapshot_id: int, db: Session = Depends(get_db)):
    """Delete a saved morning snapshot from the dashboard list."""
    service = ReportService(db)
    deleted = service.delete_morning_snapshot(snapshot_id)
    if not deleted:
        raise HTTPException(status_code=404, detail="Saved morning snapshot not found.")
    return {"message": "Saved morning snapshot deleted."}


@router.get("/history")
def get_history(skip: int = 0, limit: int = 50, db: Session = Depends(get_db)):
    reports = (
        db.query(Report)
        .order_by(Report.created_at.desc())
        .offset(skip)
        .limit(limit)
        .all()
    )
    return [
        {
            "id": r.id,
            "file_name": r.file_name,
            "sprint": r.sprint,
            "sprint_id": r.sprint_id,
            "sprint_start": r.sprint_start,
            "sprint_end": r.sprint_end,
            "day1_fixed_scope": r.day1_fixed_scope,
            "report_date": r.report_date,
            "snapshot_type": r.snapshot_type,
            "total_scope": r.total_scope,
            "completed_sp": r.completed_sp,
            "completion_percentage": r.completion_percentage,
            "created_at": r.created_at,
        }
        for r in reports
    ]


@router.get("/{report_id}")
def get_report(report_id: int, db: Session = Depends(get_db)):
    report = db.query(Report).filter(Report.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="Report not found.")
    return report


@router.get("/{report_id}/download")
def download_report(report_id: int, db: Session = Depends(get_db)):
    report = db.query(Report).filter(Report.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="Report not found.")
    if not report.file_path or not os.path.exists(report.file_path):
        raise HTTPException(status_code=404, detail="Excel file not found on disk.")
    return FileResponse(
        path=report.file_path,
        filename=report.file_name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@router.get("/{report_id}/preview", response_class=HTMLResponse)
def preview_report(report_id: int, db: Session = Depends(get_db)):
    """Render the generated Excel report as a browser-viewable HTML preview."""
    report = db.query(Report).filter(Report.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="Report not found.")
    if not report.file_path or not os.path.exists(report.file_path):
        raise HTTPException(status_code=404, detail="Excel file not found on disk.")

    from openpyxl import load_workbook

    # Keep formula source available for the download, but render the simple
    # report formulas as values in the browser preview.
    workbook = load_workbook(report.file_path, read_only=True, data_only=False)

    def display_value(worksheet, cell):
        value = cell.value
        if not (isinstance(value, str) and value.startswith("=")):
            return value

        expression = value[1:].strip()
        if expression.startswith("IFERROR(") and expression.endswith(',"N/A")'):
            expression = expression[len("IFERROR("):-len(',"N/A")')]
        evaluated = evaluate_expression(worksheet, expression)
        if (
            isinstance(evaluated, (int, float))
            and "%" in (cell.number_format or "")
        ):
            return f"{evaluated * 100:.2f}%"
        return evaluated

    def evaluate_expression(worksheet, expression):
        if "/" in expression and not expression.startswith("SUM("):
            try:
                left, right = expression.split("/", 1)
                numerator = float(display_value(worksheet, worksheet[left]) or 0)
                denominator = float(display_value(worksheet, worksheet[right]) or 0)
                return numerator / denominator if denominator else "N/A"
            except (ValueError, ZeroDivisionError):
                return "N/A"
        if expression.startswith("SUM(") and expression.endswith(")"):
            try:
                total = 0.0
                for row in worksheet[expression[4:-1]]:
                    values = row if isinstance(row, tuple) else (row,)
                    total += sum(float(display_value(worksheet, item) or 0) for item in values)
                return total
            except (TypeError, ValueError, KeyError):
                return "N/A"
        if "-" in expression:
            try:
                left, right = expression.split("-", 1)
                return float(display_value(worksheet, worksheet[left]) or 0) - float(
                    display_value(worksheet, worksheet[right]) or 0
                )
            except (TypeError, ValueError, KeyError):
                return "N/A"
        return value

    tables = []
    for worksheet in workbook.worksheets:
        rows = []
        for row in worksheet.iter_rows():
            cells = "".join(
                f"<td>{escape('' if display_value(worksheet, cell) is None else str(display_value(worksheet, cell)))}</td>"
                for cell in row
            )
            if cells.replace("<td></td>", ""):
                rows.append(f"<tr>{cells}</tr>")
        tables.append(
            f"<section><h2>{escape(worksheet.title)}</h2>"
            f"<div class='table-wrap'><table>{''.join(rows)}</table></div></section>"
        )
    workbook.close()
    title = escape(report.file_name or "Sprint Report")
    return HTMLResponse(
        f"""<!doctype html><html><head><meta charset="utf-8"><title>{title}</title>
<style>
body{{font-family:Arial,sans-serif;background:#f3f6fa;color:#172033;margin:0;padding:24px}}
main{{max-width:1400px;margin:auto;background:#fff;padding:24px;border-radius:12px;box-shadow:0 2px 12px #0001}}
h1{{margin-top:0;color:#1f3864}} h2{{color:#2e75b6;margin-top:28px}}
.table-wrap{{overflow:auto;border:1px solid #d8dee8;border-radius:8px}}
table{{border-collapse:collapse;min-width:900px;width:100%}} td{{border:1px solid #d8dee8;padding:8px;white-space:pre-wrap}}
tr:first-child td{{font-weight:700;background:#1f3864;color:#fff}}
</style></head><body><main><h1>{title}</h1>{''.join(tables)}
<p><a href="/api/reports/{report_id}/download">Download original Excel file</a></p>
</main></body></html>"""
    )


@router.delete("/{report_id}")
def delete_report(report_id: int, db: Session = Depends(get_db)):
    report = db.query(Report).filter(Report.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="Report not found.")

    file_path = report.file_path
    if file_path:
        try:
            _delete_file_with_retry(file_path)
        except FileNotFoundError:
            pass
        except PermissionError as exc:
            raise HTTPException(
                status_code=409,
                detail="The report file is currently in use and cannot be deleted. Please close the file and try again.",
            ) from exc

    db.delete(report)
    db.commit()
    return {"message": "Report deleted."}
