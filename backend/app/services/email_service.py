from __future__ import annotations

import json
import os
import re
import time
from datetime import datetime, timezone
from email.utils import parseaddr
from typing import Iterable, Optional

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.config import settings
from app.integrations.email.gmail import GmailEmailProvider
from app.models.email_delivery import EmailDelivery
from app.models.report import Report
from app.models.settings_model import AppSetting

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def _valid_addresses(addresses: Iterable[str]) -> list[str]:
    result = []
    for value in addresses:
        address = parseaddr(value)[1].strip()
        if not EMAIL_RE.match(address):
            raise ValueError(f"Invalid email recipient: {value}")
        result.append(address)
    return result


class EmailAutomationService:
    def __init__(self, db: Session):
        self.db = db
        self.provider = GmailEmailProvider()

    def settings(self) -> dict:
        return {
            "sender_email": self._get("email_sender", settings.GMAIL_USER),
            "recipients": self._get_json("email_recipients", []),
            "cc": self._get_json("email_cc", []),
            "subject": self._get("email_subject", ""),
            "report_type": self._get("email_report_type", "EOD"),
            "morning_time": self._get("email_morning_time", "09:30"),
            "eod_time": self._get("email_eod_time", "18:30"),
            "timezone": self._get("email_timezone", "Asia/Kolkata"),
            "days": self._get_json("email_days", [0, 1, 2, 3, 4]),
            "enabled": bool(self._get_json("email_enabled", False)),
            "jira_csv_path": self._get("jira_csv_path", settings.JIRA_CSV_PATH),
        }

    def update_settings(self, values: dict) -> dict:
        allowed = {
            "sender_email": "email_sender",
            "recipients": "email_recipients",
            "cc": "email_cc",
            "subject": "email_subject",
            "report_type": "email_report_type",
            "morning_time": "email_morning_time",
            "eod_time": "email_eod_time",
            "timezone": "email_timezone",
            "days": "email_days",
            "enabled": "email_enabled",
            "jira_csv_path": "jira_csv_path",
        }
        for field, key in allowed.items():
            if field in values:
                value = values[field]
                if field in {"recipients", "cc", "days", "enabled"}:
                    value = json.dumps(value)
                self._set(key, value)
        return self.settings()

    def validate_configuration(self) -> dict:
        try:
            self.provider.validate_configuration()
            configured = True
            error = None
        except ValueError as exc:
            configured = False
            error = str(exc)
        return {"configured": configured, "sender_email": self.settings()["sender_email"], "error": error}

    def send_report(self, report: Report, *, test: bool = False, recipient: Optional[str] = None) -> EmailDelivery:
        config = self.settings()
        recipients = _valid_addresses([recipient] if recipient else config["recipients"])
        cc = _valid_addresses(config["cc"])
        if not recipients:
            raise ValueError("Configure at least one email recipient.")
        if not report.file_path or not os.path.isfile(report.file_path):
            raise FileNotFoundError("Report Excel attachment does not exist.")
        workbook = load_workbook(report.file_path, read_only=False, data_only=False)
        if "Sprint Report" not in workbook.sheetnames:
            raise ValueError("Generated workbook is missing the Sprint Report sheet.")
        workbook.close()

        report_type = "TEST" if test else report.snapshot_type
        key = f"{report_type}:{report.sprint_id or report.sprint}:{report.report_date}"
        if test:
            key = f"TEST:{report.id}:{datetime.now(timezone.utc).isoformat()}"
        existing = self.db.query(EmailDelivery).filter_by(idempotency_key=key).first()
        if existing and existing.status == "SENT":
            return existing
        delivery = existing or EmailDelivery(
            idempotency_key=key,
            report_type=report_type,
            sprint=report.sprint or "",
            sprint_id=report.sprint_id,
            report_date=report.report_date or "",
            recipients=json.dumps(recipients + cc),
            attachment_filename=report.file_name,
            status="PENDING",
        )
        delivery.status = "GENERATED"
        delivery.error_message = None
        self.db.add(delivery)
        self.db.commit()
        subject = config["subject"] or f"DC-AI {report.sprint} — {report.snapshot_type} Report | {report.report_date}"
        if test:
            subject = f"[Test] {subject}"
        body = (
            f"Hi Team,\n\nPlease find attached the DC-AI {report.sprint} development "
            f"progress report for {report.report_date}.\n\nRegards,\n"
            "Jira Sprint Reporting Automation"
        )
        last_error = None
        for attempt in range(3):
            try:
                self.provider.send(
                    to=recipients,
                    cc=cc,
                    subject=subject,
                    body=body,
                    attachment_path=report.file_path,
                    attachment_filename=report.file_name,
                )
                last_error = None
                break
            except Exception as exc:
                last_error = exc
                delivery.retry_count += 1
                delivery.error_message = str(exc)
                self.db.commit()
                if attempt < 2:
                    time.sleep(0.2 * (attempt + 1))
        if last_error is not None:
            delivery.status = "FAILED"
            self.db.commit()
            raise last_error
        delivery.status = "SENT"
        delivery.sent_at = datetime.now(timezone.utc)
        self.db.commit()
        return delivery

    def history(self, limit: int = 100) -> list[dict]:
        rows = self.db.query(EmailDelivery).order_by(EmailDelivery.generated_at.desc()).limit(limit).all()
        return [
            {
                "id": row.id, "report_type": row.report_type, "sprint": row.sprint,
                "report_date": row.report_date, "generated_at": row.generated_at,
                "recipients": json.loads(row.recipients), "attachment_filename": row.attachment_filename,
                "sent_at": row.sent_at, "status": row.status, "retry_count": row.retry_count,
                "error_message": row.error_message,
            }
            for row in rows
        ]

    def _get(self, key: str, default=None):
        row = self.db.query(AppSetting).filter_by(key=key).first()
        return default if not row else row.value

    def _get_json(self, key: str, default):
        value = self._get(key, None)
        if value is None:
            return default
        try:
            return json.loads(value)
        except (TypeError, ValueError):
            return default

    def _set(self, key: str, value) -> None:
        row = self.db.query(AppSetting).filter_by(key=key).first()
        serialized = value if isinstance(value, str) else json.dumps(value)
        if row:
            row.value = serialized
        else:
            self.db.add(AppSetting(key=key, value=serialized))
        self.db.commit()
