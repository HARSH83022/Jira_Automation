import json

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.config import settings
from app.database import Base
from app.integrations.email.gmail import GmailEmailProvider
from app.models.report import Report
from app.services.email_service import EmailAutomationService


@pytest.fixture
def db():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    yield session
    session.close()


def _report(tmp_path):
    path = tmp_path / "DC_AI_Sprint_2_Report_2026-09-14.xlsx"
    wb = Workbook()
    wb.active.title = "Sprint Report"
    wb.save(path)
    return Report(
        file_name=path.name,
        sprint="Sprint 2",
        sprint_id="42",
        report_date="14/09/2026",
        snapshot_type="EOD",
        file_path=str(path),
    )


def test_gmail_configuration_reads_environment_without_exposing_secret(monkeypatch):
    monkeypatch.setattr(settings, "GMAIL_USER", "sender@example.com")
    monkeypatch.setattr(settings, "GMAIL_APP_PASSWORD", "secret-not-printed")
    provider = GmailEmailProvider()
    provider.validate_configuration()
    assert provider.user == "sender@example.com"


def test_gmail_configuration_fails_without_credentials(monkeypatch):
    monkeypatch.setattr(settings, "GMAIL_USER", None)
    monkeypatch.setattr(settings, "GMAIL_APP_PASSWORD", None)
    with pytest.raises(ValueError, match="GMAIL_USER"):
        GmailEmailProvider().validate_configuration()


def test_attachment_validation_and_delivery_history(db, tmp_path, monkeypatch):
    report = _report(tmp_path)
    db.add(report)
    db.commit()
    service = EmailAutomationService(db)

    class FakeProvider:
        def send(self, **kwargs):
            assert kwargs["attachment_filename"] == report.file_name
            assert load_workbook(kwargs["attachment_path"]).sheetnames == ["Sprint Report"]

    monkeypatch.setattr(service, "provider", FakeProvider())
    delivery = service.send_report(report, recipient="team@example.com")
    assert delivery.status == "SENT"
    assert service.history()[0]["attachment_filename"] == report.file_name


def test_duplicate_send_prevention(db, tmp_path, monkeypatch):
    report = _report(tmp_path)
    db.add(report)
    db.commit()
    service = EmailAutomationService(db)
    calls = []

    class FakeProvider:
        def send(self, **kwargs):
            calls.append(kwargs)

    monkeypatch.setattr(service, "provider", FakeProvider())
    first = service.send_report(report, recipient="team@example.com")
    second = service.send_report(report, recipient="team@example.com")
    assert first.id == second.id
    assert len(calls) == 1


def test_failed_send_records_retry_and_error(db, tmp_path, monkeypatch):
    report = _report(tmp_path)
    db.add(report)
    db.commit()
    service = EmailAutomationService(db)

    class FailingProvider:
        def send(self, **kwargs):
            raise TimeoutError("temporary SMTP failure")

    monkeypatch.setattr(service, "provider", FailingProvider())
    with pytest.raises(TimeoutError):
        service.send_report(report, recipient="team@example.com")
    delivery = service.history()[0]
    assert delivery["status"] == "FAILED"
    assert delivery["retry_count"] == 3
    assert "temporary SMTP failure" in delivery["error_message"]
