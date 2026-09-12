"""
Unit tests for the calculation engine.
All numbers must be deterministic — AI never involved.
"""
import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.calculations.engine import SprintCalculationEngine
from app.calculations.morning_eod import compute_movement
from app.database import Base
from app.services.excel_service import ExcelService
from app.services.report_service import ReportService, _morning_only_movement

COMPLETED_STATUSES = ["Done/ Live", "Ready for QA", "QA Testing in Progress"]
TEAM = ["Raj Shinde", "Sunny Shankar", "Sriniwas Chamreddy", "Dinesh Babu", "Ayush Srivastava"]


def engine():
    return SprintCalculationEngine(TEAM, COMPLETED_STATUSES)


# ── Basic completion % ────────────────────────────────────────────────────────

def test_basic_completion():
    rows = [
        {"Issue key": "X-1", "Summary": "Story 1", "Status": "Done/ Live",
         "Custom field (Developer 1)": "Raj Shinde", "Custom field (Dev Owner 1 SP)": "10"},
        {"Issue key": "X-2", "Summary": "Story 2", "Status": "To Do",
         "Custom field (Developer 1)": "Raj Shinde", "Custom field (Dev Owner 1 SP)": "15"},
    ]
    result = engine().calculate(rows, list(rows[0].keys()))
    stats = result.developer_stats["Raj Shinde"]
    assert stats.assigned_sp == 25.0
    assert stats.completed_sp == 10.0
    assert stats.completion_pct == 40.0
    assert stats.remaining_sp == 15.0


def test_completion_40_percent():
    """Assigned=100, Completed=40 → 40%"""
    rows = [
        {"Issue key": f"X-{i}", "Summary": f"S{i}", "Status": "Done/ Live",
         "Custom field (Developer 1)": "Raj Shinde", "Custom field (Dev Owner 1 SP)": "8"}
        for i in range(5)
    ] + [
        {"Issue key": f"X-{i+5}", "Summary": f"S{i+5}", "Status": "To Do",
         "Custom field (Developer 1)": "Raj Shinde", "Custom field (Dev Owner 1 SP)": "12"}
        for i in range(5)
    ]
    result = engine().calculate(rows, list(rows[0].keys()))
    stats = result.developer_stats["Raj Shinde"]
    assert stats.assigned_sp == 100.0
    assert stats.completed_sp == 40.0
    assert stats.completion_pct == 40.0


def test_zero_sp_gives_zero_pct():
    """Assigned=0 → completion % = 0%, no division by zero."""
    result = engine().calculate([], [])
    assert result.overall_completion_pct == 0.0


def test_no_division_by_zero():
    rows = [
        {"Issue key": "X-1", "Summary": "Story", "Status": "To Do",
         "Assignee": "Unknown Person", "Custom field (Story point estimate)": "5"},
    ]
    result = engine().calculate(rows, list(rows[0].keys()))
    assert result.overall_completion_pct == 0.0


# ── Missing SP ────────────────────────────────────────────────────────────────

def test_missing_sp_not_counted():
    rows = [
        {"Issue key": "X-1", "Summary": "No SP story", "Status": "Done/ Live",
         "Custom field (Developer 1)": "Raj Shinde", "Custom field (Dev Owner 1 SP)": None},
    ]
    result = engine().calculate(rows, list(rows[0].keys()))
    assert result.stories_without_dev == 1  # no valid sp → falls to no-allocation


def test_missing_developer_counted():
    rows = [
        {"Issue key": "X-1", "Summary": "Story", "Status": "Done/ Live",
         "Custom field (Dev Owner 1 SP)": "5"},
    ]
    result = engine().calculate(rows, list(rows[0].keys()))
    assert result.stories_without_dev == 1


# ── Duplicate allocation ──────────────────────────────────────────────────────

def test_no_double_count_same_dev():
    rows = [
        {
            "Issue key": "X-1", "Summary": "Story", "Status": "Done/ Live",
            "Custom field (Developer 1)": "Raj Shinde", "Custom field (Dev Owner 1 SP)": "10",
            "Custom field (Developer 2)": "Raj Shinde", "Custom field (Dev Owner 2 SP)": "10",
        }
    ]
    result = engine().calculate(rows, list(rows[0].keys()))
    stats = result.developer_stats["Raj Shinde"]
    assert stats.assigned_sp == 10.0  # not 20 — duplicate skipped
    assert len([w for w in result.warnings if "Duplicate" in w]) == 1


# ── Two developers on same story ──────────────────────────────────────────────

def test_two_devs_one_story():
    rows = [
        {
            "Issue key": "X-1", "Summary": "Story", "Status": "Done/ Live",
            "Custom field (Developer 1)": "Raj Shinde", "Custom field (Dev Owner 1 SP)": "8",
            "Custom field (Developer 2)": "Sunny Shankar", "Custom field (Dev Owner 2 SP)": "4",
        }
    ]
    result = engine().calculate(rows, list(rows[0].keys()))
    assert result.developer_stats["Raj Shinde"].assigned_sp == 8.0
    assert result.developer_stats["Sunny Shankar"].assigned_sp == 4.0
    assert result.total_scope == 12.0


def test_dynamic_developer_columns_and_scope_fallback():
    """Jira exports may contain more than two allocation columns."""
    rows = [{
        "Issue Key": "X-1", "Summary": "Story", "Status": "Done/ Live",
        "Developer 1": "Raj Shinde", "Dev Owner 1 SP": "3",
        "Developer 2": "Sunny Shankar", "Dev Owner 2 SP": "4",
        "Developer 3": "Sriniwas Chamreddy", "Dev Owner 3 SP": "5",
    }]
    result = engine().calculate(rows, list(rows[0]))
    assert result.total_scope == 12.0
    assert result.total_completed_sp == 12.0
    assert result.developer_stats["Sriniwas Chamreddy"].assigned_sp == 5.0


# ── Assignee fallback ─────────────────────────────────────────────────────────

def test_assignee_fallback():
    rows = [
        {"Issue key": "X-1", "Summary": "Story", "Status": "Done/ Live",
         "Assignee": "Sunny Shankar",
         "Custom field (Story point estimate)": "6"},
    ]
    result = engine().calculate(rows, list(rows[0].keys()))
    assert "Sunny Shankar" in result.developer_stats
    assert result.developer_stats["Sunny Shankar"].assigned_sp == 6.0
    assert result.developer_stats["Sunny Shankar"].completed_sp == 6.0


# ── Unknown developer ─────────────────────────────────────────────────────────

def test_unknown_dev_not_in_stats():
    rows = [
        {"Issue key": "X-1", "Summary": "Story", "Status": "Done/ Live",
         "Custom field (Developer 1)": "John Unknown", "Custom field (Dev Owner 1 SP)": "5"},
    ]
    result = engine().calculate(rows, list(rows[0].keys()))
    assert "John Unknown" not in result.developer_stats
    assert any("not in team" in w for w in result.warnings)


# ── Overall sprint totals ─────────────────────────────────────────────────────

def test_overall_totals():
    rows = [
        {"Issue key": "X-1", "Summary": "S1", "Status": "Done/ Live",
         "Custom field (Developer 1)": "Raj Shinde", "Custom field (Dev Owner 1 SP)": "20"},
        {"Issue key": "X-2", "Summary": "S2", "Status": "In Progress",
         "Custom field (Developer 1)": "Sunny Shankar", "Custom field (Dev Owner 1 SP)": "30"},
    ]
    result = engine().calculate(rows, list(rows[0].keys()))
    assert result.total_scope == 50.0
    assert result.total_completed_sp == 20.0
    assert result.total_remaining_sp == 30.0
    assert result.overall_completion_pct == 40.0


def test_overall_scope_is_independent_from_developer_scope():
    rows = [
        {
            "Issue key": "X-1", "Summary": "Team story", "Status": "Done/ Live",
            "Assignee": "Raj Shinde",
            "Custom field (Developer 1)": "Raj Shinde",
            "Custom field (Dev Owner 1 SP)": "8",
            "Custom field (Story point estimate)": "13",
            "Sprint": "Sprint 7",
        },
        {
            "Issue key": "X-2", "Summary": "Unallocated sprint story", "Status": "To Do",
            "Assignee": "Product Owner",
            "Custom field (Story point estimate)": "21",
            "Sprint": "Sprint 7",
        },
    ]
    result = engine().calculate(rows, list(rows[0]))
    assert result.total_scope == 34.0
    assert result.total_completed_sp == 13.0
    assert result.developer_stats["Raj Shinde"].assigned_sp == 8.0


# ── Story counts ──────────────────────────────────────────────────────────────

def test_story_counts():
    rows = [
        {"Issue key": "X-1", "Summary": "S1", "Status": "Done/ Live",
         "Custom field (Developer 1)": "Raj Shinde", "Custom field (Dev Owner 1 SP)": "5"},
        {"Issue key": "X-2", "Summary": "S2", "Status": "Ready for QA",
         "Custom field (Developer 1)": "Raj Shinde", "Custom field (Dev Owner 1 SP)": "3"},
        {"Issue key": "X-3", "Summary": "S3", "Status": "To Do",
         "Custom field (Developer 1)": "Raj Shinde", "Custom field (Dev Owner 1 SP)": "2"},
    ]
    result = engine().calculate(rows, list(rows[0].keys()))
    assert result.total_stories == 3
    assert result.completed_stories == 2
    assert result.open_stories == 1


def test_dynamic_owner_three_and_owner_source_labels():
    rows = [{
        "Issue key": "X-1", "Summary": "S1", "Status": "Ready for QA",
        "Developer Owner 3": "Raj Shinde", "Developer 3 SP": "4",
    }]
    result = engine().calculate(rows, list(rows[0]))
    assert result.developer_stats["Raj Shinde"].assigned_sp == 4
    assert result.allocations[0].source == "OWNER_3"


def test_unknown_owner_does_not_fallback_to_assignee():
    rows = [{
        "Issue key": "X-1", "Summary": "S1", "Status": "To Do",
        "Developer Owner 1": "Unknown Person", "Dev Owner 1 SP": "4",
        "Assignee": "Raj Shinde", "Story Points": "4",
    }]
    result = engine().calculate(rows, list(rows[0]))
    assert "Raj Shinde" not in result.developer_stats
    assert result.stories_without_dev == 1


def test_baseline_completion_and_raw_movement():
    morning_rows = [{
        "Issue key": "X-1", "Summary": "S1", "Status": "Done/ Live",
        "Assignee": "Raj Shinde", "Story Points": "71.25",
    }]
    eod_rows = morning_rows + [{
        "Issue key": "X-2", "Summary": "S2", "Status": "Done/ Live",
        "Assignee": "Raj Shinde", "Story Points": "1",
    }]
    calc = engine()
    morning = calc.calculate(morning_rows, list(morning_rows[0]))
    eod = calc.calculate(eod_rows, list(eod_rows[0]))
    movement = compute_movement(morning, eod, day1_fixed_scope=197.6)
    assert movement.morning_pct == 100.0
    assert movement.eod_pct == 100.0
    assert movement.daily_movement == 0.0


def test_scope_change_uses_day1_fixed_scope_not_morning_delta():
    morning_rows = [{
        "Issue key": "M-1", "Summary": "Morning story", "Status": "Done/ Live",
        "Assignee": "Raj Shinde", "Story Points": "50",
    }]
    eod_rows = [{
        "Issue key": "E-1", "Summary": "EOD story", "Status": "Done/ Live",
        "Assignee": "Raj Shinde", "Story Points": "60",
    }]
    calc = engine()
    morning = calc.calculate(morning_rows, list(morning_rows[0]))
    eod = calc.calculate(eod_rows, list(eod_rows[0]))
    movement = compute_movement(morning, eod, day1_fixed_scope=40)
    assert movement.scope_change == 20.0
    assert movement.daily_movement == 0.0


def test_morning_only_completion_uses_current_scope_denominator():
    rows = [{
        "Issue key": "X-1", "Summary": "Morning story", "Status": "Done/ Live",
        "Assignee": "Raj Shinde", "Story Points": "8",
    }]
    calc = engine()
    morning = calc.calculate(rows, list(rows[0]))
    movement = _morning_only_movement(morning, TEAM, day1_fixed_scope=2)
    assert movement.morning_pct == 100.0
    assert movement.morning_total_scope == 8.0


def test_data_quality_audit_records_include_required_fields():
    rows = [{
        "Issue key": "X-99",
        "Summary": "Bad data",
        "Status": "To Do",
        "Assignee": "Raj Shinde",
        "Developer Owner 1": "Raj Shinde",
        "Dev Owner 1 SP": "-3",
    }]
    result = engine().calculate(rows, list(rows[0].keys()))
    assert result.data_quality_issues
    issue = result.data_quality_issues[0]
    assert issue["issue_key"] == "X-99"
    assert issue["status"] == "To Do"
    assert issue["assignee"] == "Raj Shinde"
    assert issue["developer_owner_1"] == "Raj Shinde"
    assert issue["sp_owner_1"] == "-3"
    assert "problem_description" in issue


def test_excel_generation_uses_reference_template_and_formula_cells(tmp_path):
    morning_rows = [{
        "Issue key": "X-1",
        "Summary": "Morning story",
        "Status": "Done/ Live",
        "Assignee": "Raj Shinde",
        "Story Points": "71.25",
    }]
    eod_rows = morning_rows + [{
        "Issue key": "X-2",
        "Summary": "EOD story",
        "Status": "Done/ Live",
        "Assignee": "Raj Shinde",
        "Story Points": "1",
    }]

    calc = engine()
    morning = calc.calculate(morning_rows, list(morning_rows[0].keys()))
    eod = calc.calculate(eod_rows, list(eod_rows[0].keys()))
    movement = compute_movement(morning, eod, day1_fixed_scope=197.6)

    output = tmp_path / "report.xlsx"
    ExcelService().generate(
        movement=movement,
        eod_result=eod,
        morning_result=morning,
        report_date="10/09/2026",
        sprint_label="Sprint 1",
        output_path=str(output),
        day1_fixed_scope=197.6,
        sprint_start="31/08/2026 14:06",
        developer_order=TEAM,
        scenario="EOD",
    )

    wb = load_workbook(output)
    ws = wb["Sprint Report"]
    assert ws["A1"].value.startswith("DC-AI Sprint 1")
    assert ws["A6"].value == "1. Overall Sprint Progress"
    assert ws["A13"].value == "2. Development Team Progress — Morning vs EOD"
    assert [
        ws.cell(row=row, column=1).value
        for row in range(1, ws.max_row + 1)
    ].count("2. Development Team Progress — Morning vs EOD") == 1
    assert ws["A7"].value == "Date"
    assert ws["F8"].value == '=IFERROR(D8/C8,"N/A")'
    assert ws["G9"].value == '=IFERROR(F9-F8,"N/A")'
    assert ws["H19"].value == '=IFERROR(G19-D19,"N/A")'


def test_sprint_metadata_detects_name_id_and_dates():
    metadata = SprintCalculationEngine._parse_sprint_metadata(
        "id=42,name=Sprint 2,startDate=2026-09-14T09:00:00.000Z,endDate=2026-09-28T18:00:00.000Z"
    )
    assert metadata.name == "Sprint 2"
    assert metadata.sprint_id == "42"
    assert metadata.start_date.startswith("2026-09-14")
    assert metadata.end_date.startswith("2026-09-28")


def test_sprint_metadata_changes_for_different_sprint_names():
    rows = [
        {"Issue key": "S-1", "Status": "Done/ Live", "Sprint": "Sprint 1",
         "Assignee": "Raj Shinde", "Story Points": "2"},
    ]
    result_one = engine().calculate(rows, list(rows[0]))
    rows[0]["Sprint"] = "Sprint 2"
    result_two = engine().calculate(rows, list(rows[0]))
    assert result_one.sprint_name == "Sprint 1"
    assert result_two.sprint_name == "Sprint 2"


def test_sprint_scoped_baselines_do_not_cross_contaminate(tmp_path):
    service = object.__new__(ReportService)
    from app.models.settings_model import AppSetting
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from app.database import Base

    db_engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(db_engine)
    service.db = sessionmaker(bind=db_engine)()

    key_one = service._sprint_key("Sprint 1", None, "2026-08-31")
    key_two = service._sprint_key("Sprint 2", None, "2026-09-14")
    service._save_sprint_baseline(key_one, 197.6)
    assert service._get_sprint_baseline(key_one) == 197.6
    assert service._get_sprint_baseline(key_two) is None
    assert service.db.query(AppSetting).filter_by(key=key_one).count() == 1


def test_dynamic_report_filename_preserves_historical_file(tmp_path, monkeypatch):
    service = object.__new__(ReportService)
    from app.config import settings
    monkeypatch.setattr(settings, "REPORT_STORAGE_PATH", str(tmp_path))
    first_name, first_path = service._excel_path("Sprint 1", "42", "12/09/2026")
    open(first_path, "wb").close()
    second_name, second_path = service._excel_path("Sprint 1", "42", "12/09/2026")
    assert first_name == "DC_AI_Sprint_1_42_Report_2026-09-12.xlsx"
    assert second_name == "DC_AI_Sprint_1_42_Report_2026-09-12_2.xlsx"
    assert first_path != second_path


def test_dynamic_title_changes_with_report_date_and_snapshot_type(tmp_path):
    rows = [{"Issue key": "S-1", "Status": "Done/ Live", "Sprint": "Sprint 2",
             "Assignee": "Raj Shinde", "Story Points": "2"}]
    result = engine().calculate(rows, list(rows[0]))
    movement = compute_movement(None, result, day1_fixed_scope=2)
    output = tmp_path / "dynamic.xlsx"
    ExcelService().generate(
        movement, result, None, "14/09/2026", "Sprint 2", str(output),
        day1_fixed_scope=2, sprint_start="2026-09-14",
        sprint_end="2026-09-28", developer_order=TEAM, scenario="EOD",
    )
    ws = load_workbook(output)["Sprint Report"]
    assert ws["A1"].value == "DC-AI Sprint 2 — EOD Report | 14/09/2026"
    assert ws["B4"].value == "2026-09-14"
    assert ws["B5"].value == "2026-09-28"


def test_delete_morning_snapshot_removes_saved_record():
    from app.models.morning_snapshot import MorningSnapshot

    db_engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(db_engine)
    db = sessionmaker(bind=db_engine)()

    snap = MorningSnapshot(
        sprint="Sprint 5",
        report_date="12/09/2026",
        total_scope=52.0,
        completed_sp=52.0,
        completion_pct=100.0,
        developer_data=[],
    )
    db.add(snap)
    db.commit()

    service = ReportService(db)
    assert service.delete_morning_snapshot(snap.id) is True
    assert db.query(MorningSnapshot).count() == 0

    db.close()


def test_sprint_baseline_is_sticky_and_not_overwritten():
    service = object.__new__(ReportService)
    from app.database import Base
    from app.models.settings_model import AppSetting

    db_engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(db_engine)
    service.db = sessionmaker(bind=db_engine)()

    key = service._sprint_key("Sprint 1", None, "2026-08-31")
    service._save_sprint_baseline(key, 140.0)
    service._save_sprint_baseline(key, 210.0)

    assert service._get_sprint_baseline(key) == 140.0
    assert service.db.query(AppSetting).filter_by(key=key).count() == 1
